from docx import Document

def create_docx(content):

    doc = Document()

    doc.add_heading(
        'AI Requirement Engineering Report',
        level=1
    )

    doc.add_paragraph(content)

    file_name = "project_documentation.docx"

    doc.save(file_name)

    return file_name